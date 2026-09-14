#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
合规指南 QC 机械校验脚本 —— compliance-guide-agent 工作流的「循环闸」。

复用 yoyo-qc-auditor 三层方法论的自动化形态：
  · Layer 3 结果校验：版式 1:1 对齐（行/列数 vs 模板）
  · Layer 2/3：模板国残留扫描（无串国 = 无静默失败/无残留）
  · Layer 3：数据完整性 / 占位符扫描（产出不缺失）
  · Layer 1：封面国名对齐（目标对齐）

输出 markdown 报告 + JSON，并返回 exit code 供工作流循环判断：
  0 = 通过（可交付）｜ 1 = 未通过（需修复重跑）
"""

import argparse
import json
import os
import sys

from pptx import Presentation

# 未替换占位符（任何出现都视为 P1）
PLACEHOLDERS = [
    "TODO", "TBD", "待填", "XX待填", "【】", "XXXX", "占位",
    "placeholder", "FIXME", "待补充", "待确认数据", "TODO:",
]

# 已知模板国专有名词（不会在他国指南合法出现）；可按 base 模板扩展
DEFAULT_LEAK = [
    "ABGB", "AngG", "AZG", "ASVG", "ASchG", "ArbVG", "ÖGK", "RWR",
    "红白红卡", "维也纳", "Urlaubsgeld", "Urlaubs", "Weihnachts",
    "Dienstzettel", "Kollektivvertrag", "Österreich",
]


def collect_text(prs):
    out = []
    for s in prs.slides:
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                out.append(sh.text_frame.text)
            if sh.has_table:
                for r in sh.table.rows:
                    for c in r.cells:
                        if c.text.strip():
                            out.append(c.text)
    return out


def collect_tables(prs):
    out = {}
    for i, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            if sh.has_table:
                tb = sh.table
                cells = [[c.text for c in row.cells] for row in tb.rows]
                out[(i, sh.name)] = (len(tb.rows), len(tb.columns), cells)
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", required=True, help="生成的 PPTX")
    ap.add_argument("--template", required=True, help="底版模板 PPTX")
    ap.add_argument("--leak-keywords", default="", help="额外残留词，逗号分隔")
    ap.add_argument("--leak-file", default="", help="残留词文件（每行一个，# 注释）")
    ap.add_argument("--country", default="", help="期望封面国名（用于对齐校验）")
    ap.add_argument("--cover-shape", default="Text 5", help="封面国名形状名")
    ap.add_argument("--fail-on", default="P0,P1", help="视为未通过的级别")
    ap.add_argument("--md", default="", help="markdown 报告输出路径")
    ap.add_argument("--json", default="", help="JSON 摘要输出路径")
    args = ap.parse_args()

    leak = list(DEFAULT_LEAK)
    if args.leak_keywords:
        leak += [k.strip() for k in args.leak_keywords.split(",") if k.strip()]
    if args.leak_file and os.path.exists(args.leak_file):
        with open(args.leak_file, encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if line and not line.startswith("#"):
                    leak.append(line)

    out_prs = Presentation(args.out)
    tpl_prs = Presentation(args.template)

    issues = []  # {level, area, detail, evidence}

    # ---------- Layer 3：版式 1:1 对齐 ----------
    out_tables = collect_tables(out_prs)
    tpl_tables = collect_tables(tpl_prs)

    if len(out_prs.slides) != len(tpl_prs.slides):
        issues.append(dict(level="P1", area="页数",
                           detail=f"输出 {len(out_prs.slides)} 页 vs 模板 {len(tpl_prs.slides)} 页",
                           evidence="slide count"))

    for key, (or_, oc, ocells) in tpl_tables.items():
        if key not in out_tables:
            issues.append(dict(level="P0", area=f"slide {key[0]} {key[1]}",
                               detail="模板有该表但输出缺失", evidence=key[1]))
            continue
        (nr, nc, ncells) = out_tables[key]
        if or_ != nr or oc != nc:
            issues.append(dict(level="P0", area=f"slide {key[0]} {key[1]}",
                               detail=f"行列不一致 模板 {or_}x{oc} vs 输出 {nr}x{nc}",
                               evidence=key[1]))
        # 模板有内容但输出为空 → 数据缺失
        for ri in range(min(or_, nr)):
            for ci in range(min(oc, nc)):
                tval = ocells[ri][ci].strip() if ri < len(ocells) and ci < len(ocells[ri]) else ""
                nval = ncells[ri][ci].strip() if ri < len(ncells) and ci < len(ncells[ri]) else ""
                if tval and not nval:
                    issues.append(dict(level="P1", area=f"slide {key[0]} {key[1]}",
                                       detail=f"单元格({ri + 1},{ci + 1}) 模板有内容但输出为空",
                                       evidence="empty cell"))

    # ---------- Layer 2/3：模板国残留扫描 ----------
    all_text = "\n".join(collect_text(out_prs))
    hits = [k for k in leak if k in all_text]
    if hits:
        issues.append(dict(level="P0", area="残留扫描",
                           detail=f"发现模板国残留标记: {hits}", evidence=", ".join(hits)))

    # ---------- Layer 3：占位符扫描 ----------
    ph = set()
    for t in collect_text(out_prs):
        for p in PLACEHOLDERS:
            if p.lower() in t.lower():
                ph.add(p)
                break
    if ph:
        issues.append(dict(level="P1", area="占位符",
                           detail=f"发现未替换占位符: {sorted(ph)}", evidence=", ".join(sorted(ph))))

    # ---------- Layer 1：封面国名对齐 ----------
    if args.country:
        cover = ""
        for s in out_prs.slides:
            for sh in s.shapes:
                if sh.name == args.cover_shape and sh.has_text_frame:
                    cover = sh.text_frame.text.strip()
                    break
            if cover:
                break
        if cover and args.country not in cover:
            issues.append(dict(level="P1", area="封面国名",
                               detail=f"封面为「{cover}」未含目标国「{args.country}」",
                               evidence=cover))

    # ---------- 评分（沿用 yoyo-qc-auditor 规则） ----------
    nP0 = sum(1 for x in issues if x["level"] == "P0")
    nP1 = sum(1 for x in issues if x["level"] == "P1")
    nP2 = sum(1 for x in issues if x["level"] == "P2")
    score = max(0, 100 - 30 * nP0 - 15 * nP1 - 5 * nP2)
    if score >= 90:
        verdict = "放行"
    elif score >= 70:
        verdict = "有条件放行"
    elif score >= 50:
        verdict = "返工"
    else:
        verdict = "阻断"
    fail_levels = [x.strip() for x in args.fail_on.split(",")]
    failed = any(any(x["level"] == lv for lv in fail_levels) for x in issues)

    # ---------- 报告 ----------
    md = []
    md.append("## 🔍 合规指南 QC 机械校验报告")
    md.append(f"**对象**：`{os.path.basename(args.out)}`")
    md.append(f"**模板**：`{os.path.basename(args.template)}`")
    if args.country:
        md.append(f"**目标国**：{args.country}")
    md.append("")
    md.append("### 校验项")
    md.append("| 级别 | 区域 | 详情 | 证据 |")
    md.append("|------|------|------|------|")
    if not issues:
        md.append("| — | — | 全部通过 | — |")
    for x in issues:
        md.append(f"| {x['level']} | {x['area']} | {x['detail']} | {x['evidence']} |")
    md.append("")
    md.append("### 📊 判定")
    md.append(f"- **评分**：{score}/100")
    md.append(f"- **结论**：{'🔴' if failed else '🟢'} {verdict}")
    md.append(f"- **P0**：{nP0} ｜ **P1**：{nP1} ｜ **P2**：{nP2}")
    md.append(f"- **循环闸**：{'FAIL（需修复重跑）' if failed else 'PASS（可交付）'}")
    report = "\n".join(md)

    if args.md:
        with open(args.md, "w", encoding="utf-8") as f:
            f.write(report + "\n")
    if args.json:
        with open(args.json, "w", encoding="utf-8") as f:
            json.dump(dict(score=score, verdict=verdict, failed=failed,
                           nP0=nP0, nP1=nP1, nP2=nP2, issues=issues),
                      f, ensure_ascii=False, indent=2)
    print(report)
    sys.exit(1 if failed else 0)


if __name__ == "__main__":
    main()
